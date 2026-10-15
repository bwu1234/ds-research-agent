"""Deterministic dataset-card profiler.

Walks the data root in sorted order and writes one Markdown card per file plus
a manifest. No LLM, no clock, no randomness: the same files and profiler
version produce byte-identical cards. Card text quotes data values, which are
untrusted; cards are evidence for the agent, never instructions.
"""

from __future__ import annotations

import ast
import csv
import datetime as dt
import hashlib
import io
import json
import re
import shutil
import sqlite3
import struct
import zipfile
from collections.abc import Iterable, Iterator, Sequence
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

import openpyxl
import yaml
from defusedxml import ElementTree

from ds_research_agent.catalogue.models import (
    CatalogueEntry,
    Manifest,
    ParseStatus,
    SkippedPath,
)
from ds_research_agent.config import CatalogueSettings

# Bump whenever card text or entry fields change for the same input; a new
# version means a new catalogue generation and a full reindex.
PROFILER_VERSION = 2

# Reader the sandbox is expected to use, by format. Package set decided in D3
# (D3 progress (b)); each was checked against every benchmark file.
READERS = {
    "csv": "pandas.read_csv",
    "json": "json / pandas.json_normalize",
    "xlsx": "pandas.read_excel (openpyxl)",
    "gpkg": "geopandas.read_file (pyogrio)",
    "npz": "numpy.load",
    "cdf": "cdflib",
    "tle": "sgp4 (two-line element sets)",
    "sp3": "fixed-format text (no maintained reader)",
    "hdr": "lxml / xml.etree (XML)",
    "html": "pandas.read_html (lxml)",
    "txt": "pandas.read_csv(sep=r'\\s+') or plain text",
    "dat": "pandas.read_csv(sep=r'\\s+') or plain text",
    "lst": "pandas.read_csv(sep=r'\\s+') or plain text",
}

# Tried in order; latin-1 decodes any byte sequence, so it is last.
ENCODINGS = ("utf-8-sig", "cp1252", "latin-1")
DELIMITERS = (",", ";", "\t", "|")
SNIFF_BYTES = 64 * 1024

_INT = re.compile(r"^[+-]?\d+$")
_FLOAT = re.compile(r"^[+-]?(\d+\.\d*|\.\d+|\d+)([eE][+-]?\d+)?$")
_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


class ProfileError(Exception):
    """The file has a supported format but could not be parsed."""


@dataclass(frozen=True)
class ColumnProfile:
    name: str
    type: str
    nulls: int
    examples: tuple[str, ...]
    minimum: str | None = None
    maximum: str | None = None


@dataclass(frozen=True)
class TableProfile:
    title: str | None
    rows: int
    sampled_rows: int
    columns: tuple[ColumnProfile, ...]
    sample: tuple[tuple[str, ...], ...]


@dataclass(frozen=True)
class Profile:
    """What a format-specific profiler learned, before rendering."""

    facts: tuple[tuple[str, str], ...]
    tables: tuple[TableProfile, ...]
    # Titled verbatim excerpts (text formats), rendered as code blocks.
    texts: tuple[tuple[str, tuple[str, ...]], ...] = ()


# --- value typing -----------------------------------------------------------


def _value_type(v: str) -> str:
    if _INT.match(v):
        return "integer"
    if _FLOAT.match(v):
        return "float"
    if v.lower() in ("true", "false"):
        return "boolean"
    if _DATE.match(v):
        return "date"
    return "string"


def _column_type(types: set[str]) -> str:
    if not types:
        return "empty"
    if len(types) == 1:
        return next(iter(types))
    if types == {"integer", "float"}:
        return "float"
    return "string"


def _profile_column(name: str, values: Sequence[str | None], s: CatalogueSettings) -> ColumnProfile:
    present = [v for v in values if v is not None and v != ""]
    ctype = _column_type({_value_type(v) for v in present})
    examples: list[str] = []
    for v in present:
        if v not in examples:
            examples.append(v)
        if len(examples) == s.max_example_values:
            break
    lo = hi = None
    if present and ctype in ("integer", "float"):
        nums = [float(v) for v in present]
        lo, hi = present[nums.index(min(nums))], present[nums.index(max(nums))]
    elif present and ctype == "date":
        lo, hi = min(present), max(present)
    return ColumnProfile(
        name=name,
        type=ctype,
        nulls=len(values) - len(present),
        examples=tuple(examples),
        minimum=lo,
        maximum=hi,
    )


def _table(
    title: str | None,
    header: Sequence[str],
    rows: Iterable[Sequence[str | None]],
    s: CatalogueSettings,
) -> TableProfile:
    sampled: list[Sequence[str | None]] = []
    total = 0
    for row in rows:
        total += 1
        if len(sampled) < s.sample_rows:
            sampled.append(row)
    columns = tuple(
        _profile_column(name, [r[i] if i < len(r) else None for r in sampled], s)
        for i, name in enumerate(header)
    )
    sample = tuple(tuple("" if v is None else v for v in r) for r in sampled[: s.card_sample_rows])
    return TableProfile(
        title=title, rows=total, sampled_rows=len(sampled), columns=columns, sample=sample
    )


# --- formats ----------------------------------------------------------------


def _decode(path: Path) -> tuple[str, str]:
    return _decode_bytes(path.read_bytes())


def _decode_bytes(raw: bytes) -> tuple[str, str]:
    for enc in ENCODINGS:
        try:
            text = raw.decode(enc)
        except UnicodeDecodeError:
            continue
        if enc == "utf-8-sig":
            enc = "utf-8 with BOM" if raw.startswith(b"\xef\xbb\xbf") else "utf-8"
        return text, enc
    raise AssertionError("latin-1 decodes every byte string")


def _delimiter(head: str) -> str:
    first = head.splitlines()[0] if head else ""
    counts = [(first.count(d), -i, d) for i, d in enumerate(DELIMITERS)]
    best = max(counts)
    return best[2] if best[0] > 0 else ","


def profile_csv(path: Path, s: CatalogueSettings) -> Profile:
    # Decodes the whole file to pick an encoding; acceptable for D0 fixtures.
    # D2 must stream large files (wildfire is about 1 GB).
    text, encoding = _decode(path)
    delim = _delimiter(text[:SNIFF_BYTES])
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delim)
    try:
        header = next(reader)
    except StopIteration:
        raise ProfileError("empty file") from None
    except csv.Error as e:
        raise ProfileError(f"csv: {e}") from e

    def rows() -> Iterator[list[str]]:
        try:
            yield from reader
        except csv.Error as e:
            raise ProfileError(f"csv line {reader.line_num}: {e}") from e

    table = _table(None, header, rows(), s)
    facts = (("Encoding", encoding), ("Delimiter", repr(delim)))
    return Profile(facts=facts, tables=(table,))


def _scalar(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, int | float | str):
        return str(v)
    return json.dumps(v, ensure_ascii=False, sort_keys=True)


def _records_table(title: str | None, records: list[Any], s: CatalogueSettings) -> TableProfile:
    keys: list[str] = []
    for r in records[: s.sample_rows]:
        if isinstance(r, dict):
            keys.extend(k for k in r if k not in keys)
    rows = ([_scalar(r.get(k)) if isinstance(r, dict) else None for k in keys] for r in records)
    return _table(title, keys, rows, s)


def profile_json(path: Path, s: CatalogueSettings) -> Profile:
    size = path.stat().st_size
    if size > s.max_json_bytes:
        raise ProfileError(f"{size} bytes exceeds max_json_bytes {s.max_json_bytes}")
    text, encoding = _decode(path)
    try:
        data = json.loads(text)
    except json.JSONDecodeError as e:
        raise ProfileError(f"json: {e}") from e

    facts: list[tuple[str, str]] = [("Encoding", encoding)]
    tables: list[TableProfile] = []
    if isinstance(data, list):
        facts.append(("Top level", f"array of {len(data)} items"))
        if any(isinstance(r, dict) for r in data):
            tables.append(_records_table(None, data, s))
    elif isinstance(data, dict):
        facts.append(("Top level", f"object with keys {', '.join(map(str, data))}"))
        for k, v in data.items():
            if isinstance(v, list) and any(isinstance(r, dict) for r in v):
                tables.append(_records_table(f"`{k}`", v, s))
            else:
                facts.append((f"`{k}`", _cell(_scalar(v) or "null", s)))
    else:
        facts.append(("Top level", type(data).__name__))
    return Profile(facts=tuple(facts), tables=tuple(tables))


def _str(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, dt.date | dt.time):  # datetime is a date
        return v.isoformat()
    if isinstance(v, bytes):  # e.g. geometry in a column not declared BLOB
        return f"<{len(v)} bytes>"
    return _scalar(v)


def profile_xlsx(path: Path, s: CatalogueSettings) -> Profile:
    """One table per sheet; the first non-empty row is the header. Row counts
    are the sheet's declared dimension, so huge sheets are not read to the end."""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 - openpyxl raises many types for bad files
        raise ProfileError(f"xlsx: {type(e).__name__}: {e}") from e
    tables: list[TableProfile] = []
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True, max_row=s.sample_rows + 50)
            header: list[str] | None = None
            body: list[list[str | None]] = []
            for row in rows:
                cells = [_str(v) for v in row]
                if header is None:
                    if any(c not in (None, "") for c in cells):
                        header = [c or f"column {i + 1}" for i, c in enumerate(cells)]
                    continue
                if len(body) < s.sample_rows:
                    body.append(cells)
            if header is None:
                continue
            t = _table(f"Sheet `{ws.title}`", header, body, s)
            declared = max((ws.max_row or 0) - 1, len(body))
            tables.append(TableProfile(t.title, declared, t.sampled_rows, t.columns, t.sample))
    finally:
        wb.close()
    facts = (("Sheets", ", ".join(f"`{ws.title}`" for ws in wb.worksheets)),)
    return Profile(facts=facts, tables=tuple(tables))


def _quote_ident(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def profile_gpkg(path: Path, s: CatalogueSettings) -> Profile:
    """A GeoPackage is SQLite: one table per ``gpkg_contents`` entry, geometry
    and other blob columns listed but not sampled. Opened read-only."""
    try:
        con = sqlite3.connect(f"file:{path}?mode=ro&immutable=1", uri=True)
    except sqlite3.Error as e:
        raise ProfileError(f"gpkg: {e}") from e
    tables: list[TableProfile] = []
    facts: list[tuple[str, str]] = []
    try:
        contents = con.execute(
            "SELECT table_name, data_type, COALESCE(description, '') FROM gpkg_contents "
            "ORDER BY table_name"
        ).fetchall()
        for name, kind, desc in contents:
            q = _quote_ident(name)
            info = con.execute(f"PRAGMA table_info({q})").fetchall()
            cols = [(c[1], (c[2] or "").upper()) for c in info]
            keep = [c for c, t in cols if t not in ("BLOB", "GEOMETRY") and "GEOM" not in t]
            skipped = [c for c, _ in cols if c not in keep]
            total = con.execute(f"SELECT COUNT(*) FROM {q}").fetchone()[0]
            sel = ", ".join(_quote_ident(c) for c in keep) or "NULL"
            rows = con.execute(f"SELECT {sel} FROM {q} LIMIT ?", (s.sample_rows,)).fetchall()
            t = _table(f"Layer `{name}` ({kind})", keep, ([_str(v) for v in r] for r in rows), s)
            tables.append(TableProfile(t.title, total, t.sampled_rows, t.columns, t.sample))
            if skipped:
                facts.append((f"`{name}` geometry/blob columns", ", ".join(skipped)))
            if desc:
                facts.append((f"`{name}` description", _cell(desc, s)))
    except sqlite3.Error as e:
        raise ProfileError(f"gpkg: {e}") from e
    finally:
        con.close()
    return Profile(facts=tuple(facts), tables=tuple(tables))


def _npy_header(f: Any) -> tuple[str, tuple[int, ...]]:
    if f.read(6) != b"\x93NUMPY":
        raise ProfileError("npz member is not a .npy array")
    major = f.read(2)[0]
    size_fmt = "<H" if major == 1 else "<I"
    (hlen,) = struct.unpack(size_fmt, f.read(struct.calcsize(size_fmt)))
    header = ast.literal_eval(f.read(hlen).decode("latin-1"))
    return str(header["descr"]), tuple(header["shape"])


def profile_npz(path: Path, s: CatalogueSettings) -> Profile:
    """Array names, dtypes, and shapes from each member's .npy header."""
    try:
        with zipfile.ZipFile(path) as z:
            facts = []
            for name in sorted(z.namelist()):
                with z.open(name) as f:
                    dtype, shape = _npy_header(f)
                facts.append(
                    (f"Array `{name.removesuffix('.npy')}`", f"dtype {dtype}, shape {shape}")
                )
    except (zipfile.BadZipFile, ValueError, SyntaxError, KeyError) as e:
        raise ProfileError(f"npz: {e}") from e
    return Profile(facts=tuple(facts), tables=())


def _head_lines(path: Path, limit: int) -> tuple[list[str], int, str]:
    """The first ``limit`` lines, the total line count, and the encoding.

    Only the first ``SNIFF_BYTES`` are decoded, cut at the last complete line
    so a multi-byte character is never split; the rest is only counted.
    """
    total, last = 0, b""
    with path.open("rb") as f:
        head = f.read(SNIFF_BYTES)
        for chunk in (head, *iter(lambda: f.read(1 << 20), b"")):
            total += chunk.count(b"\n")
            last = chunk[-1:]
    if last not in (b"", b"\n"):
        total += 1  # a final line without a newline
    cut = head.rfind(b"\n") + 1 if len(head) == SNIFF_BYTES and b"\n" in head else len(head)
    text, enc = _decode_bytes(head[:cut])
    return text.splitlines()[:limit], total, enc


def profile_text(path: Path, s: CatalogueSettings) -> Profile:
    lines, total, enc = _head_lines(path, s.text_preview_lines)
    if not lines:
        raise ProfileError("empty file")
    shown = tuple(_clip_line(x, s) for x in lines)
    return Profile(
        facts=(("Encoding", enc), ("Lines", str(total))),
        tables=(),
        texts=((f"First {len(shown)} lines", shown),),
    )


def profile_xml(path: Path, s: CatalogueSettings) -> Profile:
    """Leaf elements as ``path: text`` lines, in document order."""
    try:
        root = ElementTree.parse(path).getroot()
    except Exception as e:  # noqa: BLE001 - defusedxml and expat raise several types
        raise ProfileError(f"xml: {type(e).__name__}: {e}") from e
    if root is None:
        raise ProfileError("xml: no root element")
    out: list[str] = []

    def walk(el: Any, prefix: str) -> None:
        tag = el.tag.rsplit("}", 1)[-1]
        here = f"{prefix}/{tag}" if prefix else tag
        kids = list(el)
        text = " ".join((el.text or "").split())
        if not kids and text:
            out.append(f"{here}: {text}")
        for k in kids:
            if len(out) >= s.text_preview_lines:
                return
            walk(k, here)

    walk(root, "")
    shown = tuple(_clip_line(x, s) for x in out[: s.text_preview_lines])
    return Profile(
        facts=(("Root element", root.tag.rsplit("}", 1)[-1]),),
        tables=(),
        texts=(("XML fields", shown),),
    )


class _VisibleText(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.lines: list[str] = []
        self.title = ""
        self._skip = 0
        self._in_title = False
        self.tables = 0

    def handle_starttag(self, tag: str, attrs: Any) -> None:
        if tag in ("script", "style"):
            self._skip += 1
        elif tag == "title":
            self._in_title = True
        elif tag == "table":
            self.tables += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style") and self._skip:
            self._skip -= 1
        elif tag == "title":
            self._in_title = False

    def handle_data(self, data: str) -> None:
        text = " ".join(data.split())
        if not text or self._skip:
            return
        if self._in_title:
            self.title += text
        else:
            self.lines.append(text)


def profile_html(path: Path, s: CatalogueSettings) -> Profile:
    text, enc = _decode(path)
    parser = _VisibleText()
    parser.feed(text)
    facts = [("Encoding", enc), ("Tables", str(parser.tables))]
    if parser.title:
        facts.insert(0, ("Title", _cell(parser.title, s)))
    shown = tuple(_clip_line(x, s) for x in parser.lines[: s.text_preview_lines])
    return Profile(facts=tuple(facts), tables=(), texts=(("Visible text", shown),))


PROFILERS = {
    "csv": profile_csv,
    "json": profile_json,
    "xlsx": profile_xlsx,
    "gpkg": profile_gpkg,
    "npz": profile_npz,
    "hdr": profile_xml,
    "html": profile_html,
    "txt": profile_text,
    "dat": profile_text,
    "lst": profile_text,
    "tle": profile_text,
    "sp3": profile_text,
    "text": profile_text,
    "py": profile_text,
}


# --- rendering --------------------------------------------------------------


def _clip(v: str, s: CatalogueSettings) -> str:
    v = " ".join(v.split())
    return v if len(v) <= s.max_cell_chars else v[: s.max_cell_chars - 1] + "…"


def _clip_line(v: str, s: CatalogueSettings) -> str:
    """A text line kept as written but bounded, and unable to close a code fence."""
    v = v.rstrip().replace("```", "'''")
    limit = 2 * s.max_cell_chars
    return v if len(v) <= limit else v[: limit - 1] + "…"


def _escape(v: str) -> str:
    return v.replace("\\", "\\\\").replace("|", "\\|")


def _cell(v: str, s: CatalogueSettings) -> str:
    return _escape(_clip(v, s))


def _render_table(t: TableProfile, s: CatalogueSettings) -> list[str]:
    out: list[str] = []
    if t.title:
        out += [f"## {t.title}", ""]
    stats = (
        f"Statistics from all {t.rows} rows."
        if t.sampled_rows == t.rows
        else f"Statistics from the first {t.sampled_rows} of {t.rows} rows."
    )
    out += [f"Rows: {t.rows}. Columns: {len(t.columns)}. {stats}", ""]
    out += ["| Column | Type | Nulls | Range | Example values |", "|---|---|---|---|---|"]
    for c in t.columns:
        rng = ""
        if c.minimum is not None and c.maximum is not None:
            rng = f"{_cell(c.minimum, s)} to {_cell(c.maximum, s)}"
        # Quote strings so a comma inside a value can't read as a separator.
        examples = ", ".join(
            _escape(json.dumps(_clip(e, s), ensure_ascii=False) if c.type == "string" else e)
            for e in c.examples
        )
        out.append(f"| {_cell(c.name, s)} | {c.type} | {c.nulls} | {rng} | {examples} |")
    if t.sample:
        out += ["", f"First {len(t.sample)} rows:", ""]
        out.append("| " + " | ".join(_cell(c.name, s) for c in t.columns) + " |")
        out.append("|" + "---|" * len(t.columns))
        for row in t.sample:
            cells = [_cell(row[i], s) if i < len(row) else "" for i in range(len(t.columns))]
            out.append("| " + " | ".join(cells) + " |")
    out.append("")
    return out


def render_card(
    front: dict[str, str],
    profile: Profile | None,
    status: ParseStatus,
    error: str | None,
    s: CatalogueSettings,
) -> str:
    # Every value is a string so rag-toolkit's equals/any_of filters match it.
    fm = yaml.safe_dump(front, sort_keys=False, allow_unicode=True, width=10_000)
    name = front["file_path"].rsplit("/", 1)[-1]
    lines = ["---", fm.rstrip("\n"), "---", f"# {name} ({front['domain']})", ""]
    lines += [f"Path: `{front['file_path']}`", f"Format: {front['format']}"]
    lines.append(f"Size: {front['size']} bytes")
    reader = READERS.get(front["format"])
    if reader:
        lines.append(f"Reader: {reader}")
    if status != "ok":
        lines.append(f"Parse status: {status}" + (f" ({error})" if error else ""))
    if profile is not None:
        lines += [f"{k}: {v}" for k, v in profile.facts]
    lines.append("")
    if profile is not None:
        for t in profile.tables:
            lines += _render_table(t, s)
        for title, body in profile.texts:
            lines += [f"## {title}", "", "```text", *body, "```", ""]
    return "\n".join(lines).rstrip("\n") + "\n"


# --- catalogue build ----------------------------------------------------------


def _sha256(data: bytes) -> str:
    return "sha256:" + hashlib.sha256(data).hexdigest()


def catalogue_id(domain: str, file_path: str) -> str:
    return f"{domain}/{hashlib.sha256(file_path.encode()).hexdigest()[:10]}"


def _walk(root: Path) -> tuple[list[Path], list[SkippedPath]]:
    files: list[Path] = []
    skipped: list[SkippedPath] = []
    for p in sorted(root.rglob("*"), key=lambda q: q.relative_to(root).as_posix()):
        rel = p.relative_to(root).as_posix()
        if any(part.startswith(".") for part in p.relative_to(root).parts):
            if p.is_file():
                skipped.append(SkippedPath(path=rel, reason="hidden"))
            continue
        if p.is_symlink():
            skipped.append(SkippedPath(path=rel, reason="symlink"))
        elif p.is_file():
            if "/" not in rel:
                skipped.append(SkippedPath(path=rel, reason="outside any domain directory"))
            elif "\n" in rel or "\r" in rel:
                skipped.append(SkippedPath(path=rel, reason="newline in path"))
            else:
                files.append(p)
    return files, skipped


def profile_file(path: Path, root: Path, s: CatalogueSettings) -> tuple[CatalogueEntry, str]:
    rel = path.relative_to(root).as_posix()
    domain = rel.split("/", 1)[0]
    fmt = path.suffix.lower().lstrip(".") or "none"
    raw = path.read_bytes()
    cid = catalogue_id(domain, rel)

    profiler = PROFILERS.get(fmt)
    profile: Profile | None = None
    status: ParseStatus = "ok"
    error: str | None = None
    if profiler is None:
        status = "unsupported"
    else:
        try:
            profile = profiler(path, s)
        except ProfileError as e:
            status, error = "error", str(e)

    front = {
        "title": rel,
        "domain": domain,
        "catalogue_id": cid,
        "file_path": rel,
        "file_hash": _sha256(raw),
        "format": fmt,
        "size": str(len(raw)),
        "card_kind": "file",
        "parse_status": status,
        "profiler_version": str(PROFILER_VERSION),
    }
    card = render_card(front, profile, status, error, s)
    entry = CatalogueEntry(
        catalogue_id=cid,
        domain=domain,
        file_path=rel,
        size=len(raw),
        file_hash=front["file_hash"],
        format=fmt,
        document_id=f"{rel}.md",
        card_hash=_sha256(card.encode()),
        profiler_version=PROFILER_VERSION,
        parse_status=status,
        parse_error=error,
    )
    return entry, card


def build_catalogue(s: CatalogueSettings) -> Manifest:
    """Profile ``s.data_root`` into ``s.cards_dir`` and write the manifest.

    Cards are written to a staging directory and swapped in only after every
    file has been profiled, so a failed build leaves the previous cards intact.
    """
    root = s.data_root.resolve()
    files, skipped = _walk(root)

    staging = s.cards_dir.with_name(s.cards_dir.name + ".staging")
    if staging.exists():
        shutil.rmtree(staging)
    entries: list[CatalogueEntry] = []
    seen: dict[str, str] = {}
    for path in files:
        entry, card = profile_file(path, root, s)
        if entry.catalogue_id in seen:
            raise RuntimeError(
                f"catalogue ID collision: {entry.file_path} and {seen[entry.catalogue_id]}"
            )
        seen[entry.catalogue_id] = entry.file_path
        out = staging / entry.document_id
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(card, encoding="utf-8")
        entries.append(entry)

    manifest = Manifest(
        profiler_version=PROFILER_VERSION, corpus=s.corpus, entries=entries, skipped=skipped
    )
    if s.cards_dir.exists():
        shutil.rmtree(s.cards_dir)
    staging.rename(s.cards_dir)
    s.manifest_path.parent.mkdir(parents=True, exist_ok=True)
    s.manifest_path.write_text(manifest.model_dump_json(indent=2) + "\n", encoding="utf-8")
    return manifest
