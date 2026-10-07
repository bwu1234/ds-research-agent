from __future__ import annotations

import asyncio
from collections.abc import Sequence
from pathlib import Path

import pytest

from ds_research_agent.config import Settings
from ds_research_agent.ledger import Ledger
from ds_research_agent.models import ChatMessage, ChatResult, ModelResponseError, ToolSpec, Usage
from eval.kramabench import report
from eval.kramabench.baselines import (
    Harness,
    ReplayClient,
    ReplayMismatch,
    build_messages,
    extract_answer,
    inline_files,
    run_batch,
)
from eval.kramabench.run import new_batch
from eval.kramabench.scoring import PROFILE
from eval.kramabench.split import split_for
from eval.kramabench.tasks import Task
from tests.kramabench_synthetic import settings_for

# --- answer extraction -----------------------------------------------------------


@pytest.mark.parametrize(
    ("content", "status", "value"),
    [
        ('Reasoning...\n{"answer": 42}', "ok", 42),
        ('```json\n{"answer": ["a", "b"]}\n```', "ok", ["a", "b"]),
        ('{"answer": 1} then revised {"answer": 2}', "ok", 2),
        ('{"answer": null}', "ok", None),
        ('{"answer": NaN}', "ok", None),  # value checked separately
        ('{"result": 3}', "no_answer_key", None),
        ("The answer is 42.", "no_json", None),
        ('{not json} {"answer": "x"}', "ok", "x"),
        ("  ", "empty", None),
        ('{"answer": __import__("os")}', "no_json", None),
    ],
)
def test_extract_answer(content: str, status: str, value: object) -> None:
    got = extract_answer(content)
    assert got.status == status
    if "NaN" not in content:
        assert got.value == value


# --- inlining --------------------------------------------------------------------


def test_inline_budget_is_water_filled_and_deterministic(tmp_path: Path) -> None:
    d = tmp_path / "legal" / "input"
    d.mkdir(parents=True)
    (d / "small.csv").write_text("a\n1\n")
    (d / "big.csv").write_text("".join(f"{i},{i * i}\n" for i in range(2000)))
    (d / "big2.csv").write_bytes("caf\xe9;1\n".encode("cp1252") * 2000)
    (d / "sheet.xlsx").write_bytes(b"PK\x03\x04corrupt")
    (d / "blob.dat").write_bytes(b"\0\1\2" * 10)
    files = [
        f"legal/input/{n}" for n in ("small.csv", "big.csv", "big2.csv", "sheet.xlsx", "blob.dat")
    ]
    a = inline_files(files, tmp_path, 1000)
    assert a == inline_files(list(reversed(files)), tmp_path, 1000)
    by = {f.path.rsplit("/", 1)[1]: f for f in a}
    assert by["small.csv"].status == "full" and by["small.csv"].shown_chars == 4
    assert by["sheet.xlsx"].status == "omitted_unreadable"
    assert by["blob.dat"].status == "omitted_binary"  # NUL bytes
    assert by["big.csv"].status == "truncated" and by["big.csv"].text.endswith("\n")
    assert by["big2.csv"].encoding == "cp1252" and "café" in by["big2.csv"].text
    assert sum(f.shown_chars for f in a) <= 1000
    # The small file's unused share went to the others (498 each before line cuts).
    assert by["big.csv"].shown_chars > 1000 // 3


def test_xlsx_sheets_render_as_csv(tmp_path: Path) -> None:
    import datetime

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    ws.append(["name", "when", "n"])
    ws.append(["a, b", datetime.date(2024, 1, 2), 1.5])
    wb.create_sheet("Notes").append(["note"])
    d = tmp_path / "bio" / "input"
    d.mkdir(parents=True)
    wb.save(d / "w.xlsx")
    (f,) = inline_files(["bio/input/w.xlsx"], tmp_path, 1000)
    assert f.status == "full" and f.encoding == "xlsx as csv"
    assert f.text == (
        '# sheet: Data\nname,when,n\n"a, b",2024-01-02T00:00:00,1.5\n# sheet: Notes\nnote\n'
    )
    (g,) = inline_files(["bio/input/w.xlsx"], tmp_path, 20)
    assert g.status == "truncated" and g.shown_chars <= 20


# --- runs, replay, and reports ---------------------------------------------------


class FakeClient:
    """Answers each task by its key: a correct answer, prose, an error, or a hang."""

    def __init__(self, tasks: dict[str, Task], behaviour: dict[str, str]) -> None:
        self.tasks, self.behaviour = tasks, behaviour
        self.current: Task | None = None

    async def chat(
        self, messages: Sequence[ChatMessage], tools: Sequence[ToolSpec] = ()
    ) -> ChatResult:
        assert self.current is not None
        how = self.behaviour.get(self.current.key, "right")
        if how == "hang":
            await asyncio.sleep(10)
        if how == "error":
            raise ModelResponseError("boom", 500)
        if how == "length":
            return ChatResult(
                message=ChatMessage(role="assistant", content="", thinking="long..."),
                usage=Usage(output_tokens=8192),
                done_reason="length",
                wall_s=1.0,
            )
        import json

        content = (
            "I am not sure."
            if how == "prose"
            else f"Thinking done.\n{json.dumps({'answer': self.current.gold.answer})}"
        )
        return ChatResult(
            message=ChatMessage(role="assistant", content=content, thinking="hmm"),
            usage=Usage(
                prompt_eval_tokens=100,
                output_tokens=20,
                prompt_eval_ns=2_000_000_000,
                eval_ns=1_000_000_000,
            ),
            done_reason="stop",
            model="fake",
            wall_s=3.0,
        )


def _harness(tmp_path: Path) -> tuple[Settings, Harness]:
    s = settings_for(tmp_path, task_timeout_s=0.2)
    split, tasks = split_for(s)
    s = s.model_copy(update={"eval": s.eval.model_copy(update={"split_sha256": split.sha256()})})
    return s, Harness(s, s.model, Ledger(s.eval.ledger_path), split, tasks)


def _run(h: Harness, condition: str, keys: tuple[str, ...], behaviour: dict[str, str]) -> str:
    batch = new_batch(
        h.settings, h.model, h.split, condition=condition, task_set="dev",
        keys=keys, repeats=1, batch_id=f"{condition}-test",
    )  # fmt: skip
    h.ledger.add_batch(batch)
    fake = FakeClient(h.tasks, behaviour)

    def client_for(t: Task, _r: int) -> FakeClient:
        fake.current = t
        return fake

    asyncio.run(run_batch(h, batch, client_for))
    return batch.batch_id


def test_failures_stay_in_the_denominator_and_replay_is_identical(tmp_path: Path) -> None:
    s, h = _harness(tmp_path)
    keys = h.split.dev
    behaviour = {keys[0]: "prose", keys[1]: "error", keys[2]: "hang", keys[3]: "length"}
    bid = _run(h, "inline", keys, behaviour)
    runs = {r.task_key: r for r in h.ledger.runs(bid)}
    assert [runs[k].stop_reason for k in keys[:4]] == [
        "malformed", "model_error", "timeout", "budget_exhausted",
    ]  # fmt: skip
    assert all(runs[k].stop_reason == "answered" for k in keys[4:])
    manifest = runs[keys[4]].input_manifest
    assert manifest["files"][0]["sha256"] and manifest["answer_type_visible"] is False
    # Gold never reaches the prompt.
    for k in keys:
        req = h.ledger.steps(runs[k].run_id)[0].request
        assert "Expected answer type" not in req["messages"][1]["content"]

    batch = h.ledger.batch(bid)
    out = report.summarise(h.ledger, batch, h.tasks, 200, 0)
    assert out["runs_planned"] == len(keys)
    assert out["answer_score"] == pytest.approx((len(keys) - 4) / len(keys), abs=1e-4)
    assert out["strict_accuracy"] == out["answer_score"]
    assert out["answer_score_ci95"][0] <= out["answer_score"] <= out["answer_score_ci95"][1]
    assert out["stop_reasons"] == {
        "malformed": 1, "model_error": 1, "timeout": 1, "budget_exhausted": 1,
        "answered": len(keys) - 4,
    }  # fmt: skip

    # Replay through recorded responses: same answers and scores, no model.
    rec = {(r.task_key, r.repeat): h.ledger.steps(r.run_id) for r in h.ledger.runs(bid)}
    replay = new_batch(
        s, h.model, h.split, condition="inline", task_set="dev", keys=keys, repeats=1,
        batch_id="replay", replay_of=batch,
    )  # fmt: skip
    h.ledger.add_batch(replay)
    asyncio.run(run_batch(h, replay, lambda t, r: ReplayClient(h.model, rec[(t.key, r)])))
    for a, b in zip(h.ledger.runs(bid), h.ledger.runs("replay"), strict=True):
        assert a.stop_reason == b.stop_reason
        aa, ba = h.ledger.answer(a.run_id), h.ledger.answer(b.run_id)
        assert aa and ba and aa.model_dump(exclude={"run_id"}) == ba.model_dump(exclude={"run_id"})
        sa, sb = h.ledger.score(a.run_id, PROFILE), h.ledger.score(b.run_id, PROFILE)
        assert sa and sb and (sa.score, sa.strict) == (sb.score, sb.strict)


def test_replay_refuses_a_changed_request(tmp_path: Path) -> None:
    s, h = _harness(tmp_path)
    keys = h.split.dev[:1]
    bid = _run(h, "no_tools", keys, {})
    steps = h.ledger.steps(h.ledger.runs(bid)[0].run_id)
    other = h.model.model_copy(update={"think": "xhigh"})
    msgs, _ = build_messages(h, "no_tools", h.tasks[keys[0]])
    with pytest.raises(ReplayMismatch):
        asyncio.run(ReplayClient(other, steps).chat(msgs))


def test_missing_runs_count_as_zero_and_paired_compare(tmp_path: Path) -> None:
    s, h = _harness(tmp_path)
    keys = h.split.dev
    base = _run(h, "no_tools", keys, {k: "prose" for k in keys})
    good = _run(h, "inline", keys, {})
    out = report.compare(h.ledger, h.ledger.batch(base), h.ledger.batch(good), h.tasks, 200, 0)
    assert out["score"]["diff"] == 1.0 and out["score"]["diff_ci95"] == [1.0, 1.0]
    # A planned task with no run row still counts in the denominator.
    partial = new_batch(
        s, h.model, h.split, condition="inline", task_set="dev", keys=keys, repeats=1,
        batch_id="partial",
    )  # fmt: skip
    h.ledger.add_batch(partial)
    summary = report.summarise(h.ledger, partial, h.tasks, 50, 0)
    assert summary["runs_recorded"] == 0
    assert summary["stop_reasons"] == {"missing": len(keys)}
    assert summary["answer_score"] == 0.0


def test_choose_think_prefers_fast_levels_within_one_task() -> None:
    def summ(think: str, score: float, wall: float) -> dict[str, object]:
        return {"task_set": "sample", "tasks": 10, "think": think,
                "answer_score": score, "wall_s": {"median": wall}}  # fmt: skip

    pick = report.choose_think(
        [
            summ("off", 0.10, 5),
            summ("low", 0.30, 40),
            summ("medium", 0.35, 90),
            summ("xhigh", 0.38, 300),
        ]  # fmt: skip
    )
    assert pick["eligible"] == ["low", "medium", "xhigh"]
    assert pick["chosen"] == "low"


def test_resume_skips_finished_runs_and_redoes_interrupted_ones(tmp_path: Path) -> None:
    from ds_research_agent.ledger import Run, utc_now
    from eval.kramabench.run import resume_problems

    s, h = _harness(tmp_path)
    keys = h.split.dev
    batch = new_batch(
        s, h.model, h.split, condition="inline", task_set="dev", keys=keys, repeats=1,
        batch_id="resumable",
    )  # fmt: skip
    h.ledger.add_batch(batch)
    fake = FakeClient(h.tasks, {})
    calls: list[str] = []

    def client_for(t: Task, _r: int) -> FakeClient:
        fake.current = t
        calls.append(t.key)
        return fake

    # Two tasks finished, then a third was interrupted mid-request.
    first = batch.model_copy(update={"task_keys": keys[:2]})
    asyncio.run(run_batch(h, first, client_for))
    h.ledger.add_run(
        Run(
            run_id=f"resumable/{keys[2]}/0", batch_id="resumable", task_key=keys[2],
            parent_task_key=h.tasks[keys[2]].parent_key, split="dev", condition="inline",
            repeat=0, input_manifest={}, access_policy="x", catalogue_generation=None,
            started=utc_now(),
        )
    )  # fmt: skip
    assert h.ledger.discard_unfinished("resumable") == [keys[2]]
    calls.clear()
    asyncio.run(run_batch(h, batch, client_for))
    assert calls == list(keys[2:])
    runs = h.ledger.runs("resumable")
    assert [r.task_key for r in runs] == list(keys)
    assert all(r.stop_reason == "answered" for r in runs)

    same = new_batch(
        s, h.model, h.split, condition="inline", task_set="dev", keys=keys, repeats=1,
        batch_id="resumable",
    )  # fmt: skip
    assert resume_problems(batch, same) == []
    changed = new_batch(
        s, h.model.model_copy(update={"think": "xhigh"}), h.split, condition="inline",
        task_set="dev", keys=keys, repeats=1, batch_id="resumable",
    )  # fmt: skip
    assert resume_problems(batch, changed) == ["model_settings"]
